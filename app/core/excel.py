"""Write the reimbursement workbook. All totals are live Excel formulas."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FONT = Font(color="0000FF")          # blue = value you may edit
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _money(cur):
    sym = {"GBP": "£", "CNY": "¥", "EUR": "€", "USD": "$"}.get(cur, "")
    return f'"{sym}"#,##0.00' if sym else "#,##0.00"


def _header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.fill, c.font, c.border = HEAD, HEAD_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def write_workbook(path, trip, claims, plan, status_of, scan_note=lambda r: ""):
    money = _money(trip.get("currency", "GBP"))
    wb = Workbook()

    # ---------------- 汇总 ----------------
    ws = wb.active
    ws.title = "汇总"
    ws["A1"] = f"出差报销汇总 — {trip['name']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"生成日期：{date.today().isoformat()}　蓝色字为可修改的输入值，其余为公式"
    ws["A2"].font = Font(italic=True, color="808080")

    _header(ws, 4, ["段", "开始日期", "结束日期", "天数"], [22, 14, 14, 16])
    periods = sorted(trip["periods"], key=lambda p: p["start"])
    first = 5
    for i, p in enumerate(periods):
        r = first + i
        ws.cell(r, 1, f"第 {i + 1} 段")
        for col, key in ((2, "start"), (3, "end")):
            c = ws.cell(r, col, date.fromisoformat(p[key]))
            c.number_format, c.font = "yyyy-mm-dd", INPUT_FONT
        ws.cell(r, 4, f"=C{r}-B{r}+1")
        for col in range(1, 5):
            ws.cell(r, col).border = BOX
    last = first + max(len(periods), 1) - 1
    prange = (f"汇总!$B${first}:$B${last}", f"汇总!$C${first}:$C${last}")

    r = last + 2
    rows = [
        ("出差总天数", f"=SUM(D{first}:D{last})", "0"),
        ("每日额度", float(trip.get("rate") or 0), money),
        ("可报销上限（天数 × 额度）", None, money),
        ("期内个人实付合计", "=SUMIFS(明细!$E:$E,明细!$H:$H,\"是\")", money),
        ("期外金额（不计入）", "=SUMIFS(明细!$E:$E,明细!$H:$H,\"否\")", money),
        ("应报销金额", None, money),
        ("超出上限自付部分", None, money),
        ("单据数（计入 / 全部）", "=COUNTIF(明细!$H:$H,\"是\")&\" / \"&COUNTA(明细!$A:$A)-1", None),
    ]
    refs = {}
    for label, val, fmt in rows:
        ws.cell(r, 1, label).font = Font(bold=True)
        refs[label] = f"B{r}"
        ws.cell(r, 2, val)
        if fmt:
            ws.cell(r, 2).number_format = fmt
        for col in (1, 2):
            ws.cell(r, col).border = BOX
        r += 1
    ws[refs["每日额度"]].font = INPUT_FONT
    ws[refs["可报销上限（天数 × 额度）"]] = f"={refs['出差总天数']}*{refs['每日额度']}"
    ws[refs["应报销金额"]] = f"=MIN({refs['期内个人实付合计']},{refs['可报销上限（天数 × 额度）']})"
    ws[refs["超出上限自付部分"]] = f"=MAX(0,{refs['期内个人实付合计']}-{refs['可报销上限（天数 × 额度）']})"
    for k in ("应报销金额",):
        ws[refs[k]].fill = KEY_FILL
        ws[refs[k]].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 28

    # ---------------- 明细（每单一行）----------------
    ws2 = wb.create_sheet("明细")
    _header(ws2, 1, ["单号", "日期", "商户", "票据张数", "个人实付金额", "文件名", "关联说明", "是否在报销期内", "备注"],
            [7, 12, 24, 9, 14, 38, 40, 14, 30])
    for i, c in enumerate(claims, 2):
        names = [plan[m["id"]] for m in c["members"]]
        link = ""
        if c["group"] and len(c["members"]) > 1:
            parts = [f"{plan[m['id']].rsplit('.', 1)[0][-2:]}: {m.get('amount') if m.get('amount') is not None else '?'}"
                     for m in c["members"]]
            link = f"关联 {len(c['members'])} 张（{'；'.join(parts)}），按个人实付计为一单"
        ws2.cell(i, 1, c["no"])
        dc = ws2.cell(i, 2, date.fromisoformat(c["date"]) if c["date"] else None)
        dc.number_format = "yyyy-mm-dd"
        ws2.cell(i, 3, c["merchant"])
        ws2.cell(i, 4, len(c["members"]))
        ac = ws2.cell(i, 5, c["amount"])
        ac.number_format = money
        ws2.cell(i, 6, "\n".join(names)).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.cell(i, 7, link).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.cell(i, 8, f'=IF(B{i}="","否",IF(COUNTIFS({prange[0]},"<="&B{i},{prange[1]},">="&B{i})>0,"是","否"))')
        ws2.cell(i, 9, c.get("note", ""))
        for col in range(1, 10):
            ws2.cell(i, col).border = BOX
        if c["amount"] is None or not c["date"]:
            for col in range(1, 10):
                ws2.cell(i, col).fill = WARN_FILL
    end = len(claims) + 1
    t = end + 1
    ws2.cell(t, 4, "合计").font = Font(bold=True)
    ws2.cell(t, 5, f"=SUM(E2:E{end})").number_format = money
    ws2.cell(t + 1, 4, "期内合计").font = Font(bold=True)
    ws2.cell(t + 1, 5, f'=SUMIFS(E2:E{end},H2:H{end},"是")').number_format = money
    ws2.freeze_panes = "A2"
    # The 汇总 sheet sums whole columns; keep the totals out of column H so they aren't double counted.

    # ---------------- 票据清单（每张图片一行）----------------
    ws3 = wb.create_sheet("票据清单")
    _header(ws3, 1, ["新文件名", "原文件名", "所属单号", "票面金额", "日期", "时间", "商户", "状态",
                     "OCR识别金额", "OCR识别日期", "人工修改", "人工核对时间", "扫描版", "识别提示"],
            [34, 46, 9, 12, 12, 8, 24, 26, 12, 12, 9, 20, 22, 50])
    claim_no = {m["id"]: c["no"] for c in claims for m in c["members"]}
    for i, r in enumerate(sorted(trip["receipts"], key=lambda r: plan[r["id"]]), 2):
        ocr = r.get("ocr", {})
        edited = (r.get("amount") != ocr.get("amount")) or (r.get("date") != ocr.get("date"))
        vals = [plan[r["id"]], r["orig_name"], claim_no.get(r["id"], "—"), r.get("amount"),
                date.fromisoformat(r["date"]) if r.get("date") else None, r.get("time") or "",
                r.get("merchant", ""), status_of(r), ocr.get("amount"),
                date.fromisoformat(ocr["date"]) if ocr.get("date") else None,
                "是" if edited else "", (r.get("verified_at") or "").replace("T", " "),
                scan_note(r), "；".join(ocr.get("warnings", []))]
        for col, v in enumerate(vals, 1):
            c = ws3.cell(i, col, v)
            c.border = BOX
            if col in (4, 9):
                c.number_format = money
            if col in (5, 10):
                c.number_format = "yyyy-mm-dd"
        if r["id"] not in claim_no:
            for col in range(1, 15):
                ws3.cell(i, col).fill = WARN_FILL
    ws3.freeze_panes = "A2"

    wb.save(path)
